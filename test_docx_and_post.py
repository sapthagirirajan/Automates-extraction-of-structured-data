from docx import Document
from pathlib import Path
import requests

fixtures = Path(__file__).parent / "fixtures"
fixtures.mkdir(exist_ok=True)

# create docx
doc = Document()
doc.add_heading('John Doe', level=1)
doc.add_paragraph('Email: john.doe@example.com')
doc.add_paragraph('Phone: +1-555-0123')

doc.add_heading('Skills', level=2)
doc.add_paragraph('Python, FastAPI, Docker, SQL')

out = fixtures / 'resume_doc.docx'
doc.save(out)
print('Created', out)

# post to service
url = 'http://127.0.0.1:8006/parse/resume'
with open(out,'rb') as f:
    r = requests.post(url, files={'file': ('resume.docx', f)})
    print('Status:', r.status_code)
    try:
        print('Response JSON:', r.json())
    except Exception:
        print('Response text:', r.text[:1000])
