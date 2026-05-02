from docx import Document
from pathlib import Path

fixtures = Path(__file__).parent / "fixtures"
fixtures.mkdir(exist_ok=True)

doc = Document()
doc.add_heading('John Doe', level=1)
doc.add_paragraph('Email: john.doe@example.com')
doc.add_paragraph('Phone: +1-555-0123')

doc.add_heading('Education', level=2)
doc.add_paragraph('B.Sc. Computer Science, Example University (2016 - 2020)')

doc.add_heading('Experience', level=2)
doc.add_paragraph('Senior Software Engineer, TechCorp (2021 - present)')
doc.add_paragraph(' - Led backend team and built APIs')

doc.add_heading('Skills', level=2)
doc.add_paragraph('Python, FastAPI, Docker, SQL')

out = fixtures / 'resume_doc.docx'
doc.save(out)
print(f'Created {out}')
